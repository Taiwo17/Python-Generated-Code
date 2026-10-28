import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill, Border, Side, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.formatting.rule import CellIsRule
from datetime import date

found = "12 Oct 2026"
indeed_all = "https://ng.indeed.com/q-devops-engineer-jobs.html"
indeed_abuja = "https://ng.indeed.com/q-devops-l-abuja-jobs.html?l=abuja&q=devops"
indeed_remote = "https://ng.indeed.com/q-remote-devops-l-lagos-jobs.html?vjk=a944a21f118b9d59"

jobs = [
    {
        "Priority": "High", "Job Title": "Junior DevOps Engineer", "Company": "Cloud Technology Hub",
        "Location": "Lagos (hybrid remote)", "Work Arrangement": "Hybrid",
        "Experience / Level": "2–4 years; junior",
        "Salary (as listed)": "₦300,000–₦400,000/month",
        "Key skills / summary": "AWS, OCI, Linux, GitHub Actions/Jenkins, Docker, Kubernetes, Helm, Terraform, Argo CD/Flux, Prometheus, Grafana, Loki, IAM/RBAC.",
        "Company information": "Cloud/DevOps hiring company; verify company profile and role details on listing.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Listing visible in search results; original post date not confirmed",
        "Application notes": "Strong match for hands-on cloud/IaC experience. Open listing and use Apply now."
    },
    {
        "Priority": "Medium", "Job Title": "IT Infrastructure/DevOps Engineer Trainee (September 2026)", "Company": "Bincom Dev Center Nigeria",
        "Location": "Nigeria (confirm exact work location)", "Work Arrangement": "Trainee / contract",
        "Experience / Level": "Entry-level; fresh graduates and NYSC candidates mentioned",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Infrastructure and DevOps trainee opportunity; listed as a 6-month to 1-year contract.",
        "Company information": "Technology training/development centre; confirm programme details with employer.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Role title references September 2026; availability must be reconfirmed",
        "Application notes": "Potential entry-level route. Confirm that applications are still open before preparing documents."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer (Junior, Intermediate, Specialist, Expert)", "Company": "GVA Partners",
        "Location": "Lagos", "Work Arrangement": "Not stated",
        "Experience / Level": "Multiple levels; exact requirements vary",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Cloud platforms, automation, CI/CD pipelines, infrastructure management, scalability and reliability.",
        "Company information": "Role appears under GVA Partners; verify employer and full job description.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results; original post date not confirmed",
        "Application notes": "Check which level best fits your experience before applying."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "Chowdeck",
        "Location": "Sangotedo, Lagos", "Work Arrangement": "Not stated",
        "Experience / Level": "Not stated",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Improve developer experience and platform efficiency; full requirements need checking.",
        "Company information": "Nigerian food, medicine and grocery delivery platform.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results; original post date not confirmed",
        "Application notes": "Review the full description for on-site expectations and cloud/tool requirements."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "Texture Science Labs Nigeria",
        "Location": "Nigeria (exact city not stated)", "Work Arrangement": "Not stated",
        "Experience / Level": "Not stated",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Work with software engineers on deployment, security, scalability and operational processes.",
        "Company information": "Company details should be verified through the employer's official channels.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results; original post date not confirmed",
        "Application notes": "Open the full listing to check required years of experience and application route."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "Helium Health",
        "Location": "Lagos", "Work Arrangement": "Not stated",
        "Experience / Level": "3–5 years",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "DevOps collaboration with product, operations, quality and software engineering teams.",
        "Company information": "Health technology company focused on healthcare systems and services.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results; original post date not confirmed",
        "Application notes": "Check whether the role is still open and review the full technical requirements."
    },
    {
        "Priority": "Low", "Job Title": "Cloud/DevOps Graduate Trainee", "Company": "Gurugeeks Royalty Limited",
        "Location": "Lagos", "Work Arrangement": "Not stated",
        "Experience / Level": "Graduate trainee / entry-level",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Hands-on training and mentorship in designing, deploying and managing cloud infrastructure.",
        "Company information": "Employer profile and programme details should be independently verified.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results; original post date not confirmed",
        "Application notes": "Potential fit if you are targeting junior/trainee opportunities."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "Teknowledge",
        "Location": "Lagos", "Work Arrangement": "Not stated",
        "Experience / Level": "Not stated",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Role supports technology for financial services and banking; see full listing for stack and experience requirements.",
        "Company information": "Technology services provider; verify the exact employing entity on the listing.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results; original post date not confirmed",
        "Application notes": "Review whether the role is DevOps-focused and confirm work arrangement."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer (100% remote)", "Company": "Tether Operations Limited",
        "Location": "Remote; listing surfaced for Lagos", "Work Arrangement": "Remote",
        "Experience / Level": "Experienced; exact years not confirmed",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "End-to-end CI/CD pipelines using GitHub; detailed requirements in full description.",
        "Company information": "Global digital-asset and financial technology company; verify role eligibility for applicants based in Nigeria.",
        "Application link": indeed_remote, "Source / job listing": indeed_remote,
        "Posting date / freshness": "Remote search page showed role; original post date not confirmed",
        "Application notes": "Confirm Nigeria eligibility, timezone expectations and official company application route."
    },
    {
        "Priority": "High", "Job Title": "Senior DevOps Engineer", "Company": "Workpedia Africa Ltd",
        "Location": "Abuja", "Work Arrangement": "In-person; listing says role only available to candidates in Abuja",
        "Experience / Level": "4–6 years; senior",
        "Salary (as listed)": "₦800,000–₦1,000,000/month",
        "Key skills / summary": "CI/CD, IaC, Kubernetes, Linux, Bash/Python, Prometheus/Grafana, reliability, security and incident response.",
        "Company information": "Recruitment listing names Workpedia Africa Ltd; verify company and role through the original posting.",
        "Application link": indeed_abuja, "Source / job listing": indeed_abuja,
        "Posting date / freshness": "Indeed Abuja page crawled 9 Oct 2026; confirm current availability",
        "Application notes": "Listing says candidates must be in Abuja and instructs applicants to send CV stating the role; confirm contact details on listing."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "ByteWorks",
        "Location": "Abuja", "Work Arrangement": "Not stated",
        "Experience / Level": "Not stated",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Role combines technical engineering and operations; detailed skills and requirements need checking.",
        "Company information": "Technology company; verify the exact role and employer details from the listing.",
        "Application link": indeed_abuja, "Source / job listing": indeed_abuja,
        "Posting date / freshness": "Visible in Abuja DevOps search results; original post date not confirmed",
        "Application notes": "Open the listing and confirm whether applicants based in Lagos can be considered."
    },
    {
        "Priority": "High", "Job Title": "Senior DevOps & Systems Administrator", "Company": "HUIOSPAY LIMITED",
        "Location": "Abuja, Nigeria", "Work Arrangement": "Remote/hybrid stated in listing",
        "Experience / Level": "Senior; 6+ years",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Multi-cloud, IaC, CI/CD, infrastructure security, scripting, Docker/Kubernetes, monitoring/logging and database systems.",
        "Company information": "Fintech/payments-related role according to the job description; verify the employer and compliance details.",
        "Application link": indeed_abuja, "Source / job listing": "https://ng.indeed.com/q-devops-l-abuja-jobs.html?l=abuja&q=devops",
        "Posting date / freshness": "Abuja results page crawled 11 Oct 2026; confirm current availability",
        "Application notes": "Listing mentions submitting CV and cover letter to career@huiospay.com; verify on the current posting before sending."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "SabiRecruit",
        "Location": "Nigeria", "Work Arrangement": "Remote, project-based contract",
        "Experience / Level": "3+ years in DevOps/SRE stated in requirements",
        "Salary (as listed)": "₦650,000/month",
        "Key skills / summary": "Linux, Docker, CI/CD, Terraform or similar IaC, Kubernetes fundamentals, cloud; evaluate infrastructure configurations and pipelines.",
        "Company information": "Recruitment/project-based listing; verify the end client, contract terms and payment details.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed search results on 12 Oct 2026; original post date not confirmed",
        "Application notes": "Confirm whether the work is project-based AI evaluation rather than a conventional in-house DevOps role."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer / Linux Administrator", "Company": "eHealth4everyone",
        "Location": "Abuja", "Work Arrangement": "Not stated",
        "Experience / Level": "Experienced; exact years not confirmed",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "DevOps and Linux administration supporting systems and customer experience.",
        "Company information": "Digital health organization; verify the specific role and current hiring status.",
        "Application link": indeed_abuja, "Source / job listing": indeed_abuja,
        "Posting date / freshness": "Visible in Abuja DevOps search results; original post date not confirmed",
        "Application notes": "Check the complete description and whether relocation or Abuja presence is required."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "Probe Compliance Africa",
        "Location": "Lagos", "Work Arrangement": "Not stated",
        "Experience / Level": "Not stated",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "Search snippet mentions Windows Server AMIs on AWS EC2 and remote development environments.",
        "Company information": "Compliance-focused technology listing; verify company details and job requirements.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results on 12 Oct 2026; original post date not confirmed",
        "Application notes": "Open the full listing to confirm role scope, work arrangement and application link."
    },
    {
        "Priority": "Medium", "Job Title": "DevOps Engineer", "Company": "GV Alliance Partners Limited",
        "Location": "Lagos", "Work Arrangement": "Full-time; 8am–5pm stated",
        "Experience / Level": "Senior; 5 years stated",
        "Salary (as listed)": "Not stated",
        "Key skills / summary": "DevOps engineering role; full technical stack not visible in the search snippet.",
        "Company information": "Employer/recruitment listing; verify the hiring organization and complete requirements.",
        "Application link": indeed_all, "Source / job listing": indeed_all,
        "Posting date / freshness": "Visible in Indeed results; original post date not confirmed",
        "Application notes": "Check required tools, salary and work location before applying."
    },
]

prompt = """Act as my DevOps job-search research assistant. Search the live web for currently open DevOps and closely related infrastructure/cloud engineering roles that I can apply for from Lagos, Nigeria.

MY LOCATION AND PREFERENCES
- I live in Lagos, Nigeria.
- Prioritise Lagos-based roles, fully remote roles that explicitly accept applicants living in Nigeria, and Abuja roles only when the listing clearly permits remote/hybrid work or is worth flagging separately.
- Search for: DevOps Engineer, Junior DevOps Engineer, Associate DevOps Engineer, Cloud Engineer, Platform Engineer, Site Reliability Engineer (SRE), Infrastructure Engineer, DevOps/Cloud Trainee, and DevSecOps Engineer.
- Include junior, entry-level, trainee, mid-level and senior roles. Do not exclude a role solely because the title differs if the duties are strongly DevOps-related.

FRESHNESS AND VERIFICATION
- Search current job boards and official company career pages, including LinkedIn Jobs, Indeed Nigeria, Jobberman, MyJobMag, Wellfound, and employer career sites where useful.
- Prefer jobs posted or refreshed within the last 30 days; label older listings clearly.
- Open each listing and check whether it appears active. Do not describe a role as open if that cannot be verified.
- Deduplicate repeated listings.
- Never invent a company, salary, posting date, requirement, contact, or application URL. Mark missing information as “Not stated” or “Not verified”.
- Prefer a direct employer application URL. If only a job-board listing is available, link to that listing and label it accordingly.
- Flag suspicious listings, unclear employers, requests for application fees, or unusual contact-only application methods.

FOR EACH JOB, COLLECT
1. Job title
2. Company name and a brief, verified company description
3. Location and work arrangement (onsite/hybrid/remote)
4. Whether a Lagos-based applicant is eligible
5. Seniority and years of experience
6. Main responsibilities and required technical skills
7. Salary/range and currency, if published
8. Employment type
9. Posting date and date checked
10. Application deadline, if stated
11. Direct application link
12. Source link
13. A short fit assessment for a candidate living in Lagos
14. Important caveats or missing details

OUTPUT
- Create a well-formatted Excel workbook (.xlsx), not just a chat list.
- Sheet 1: “Job Tracker”, with one row per unique job, filters, frozen header row, wrapped text, readable column widths, clickable application/source links, and status dropdowns (Not Applied, To Apply, Applied, Interview, Rejected, Closed).
- Sheet 2: “Summary”, with total roles, counts by location/work arrangement/seniority, and a shortlist of the best matches for a Lagos-based applicant.
- Sheet 3: “Search Log”, listing search sources, search date, filters used, and verification limitations.
- Use a clean professional style and highlight roles that are remote-eligible or particularly relevant to Lagos.
- Sort by fit and freshness, not by search-engine ranking.
- In the final response, summarise the strongest matches, explain any limitations, and link to the Excel file.
- Clearly state the date and time the search was performed. This is a snapshot, not a guarantee that a role remains open."""

with pd.ExcelWriter("DevOps_Jobs_Nigeria_Lagos_12_Oct_2026.xlsx", engine="openpyxl") as writer:
    pd.DataFrame(jobs).to_excel(writer, sheet_name="Job Tracker", index=False)
    pd.DataFrame([
        ["Search date", found],
        ["Candidate location", "Lagos, Nigeria"],
        ["Scope", "Lagos, Abuja and remote roles accessible from Nigeria"],
        ["Roles found in this snapshot", len(jobs)],
        ["Best initial leads", "Cloud Technology Hub (Junior DevOps); Tether (remote); SabiRecruit (remote contract); Probe Compliance Africa (Lagos)"],
        ["Abuja roles to assess carefully", "Workpedia Africa Ltd states Abuja-only/in-person; HUIOSPAY listing mentions remote/hybrid—confirm directly."],
        ["Important limitation", "Most source pages expose job-board search/listing pages rather than a stable direct employer application URL. Click through and confirm active status before applying."],
        ["Freshness caveat", "Search result visibility is not proof that a job is still open. Original posting dates were not consistently available."],
        ["Improved reusable prompt", prompt]
    ], columns=["Item", "Details"]).to_excel(writer, sheet_name="Summary", index=False)
    pd.DataFrame([
        ["Indeed Nigeria – DevOps roles", found, "DevOps Engineer, Junior DevOps, Cloud/DevOps trainee", "https://ng.indeed.com/q-devops-engineer-jobs.html", "Search results contained multiple Lagos/Nigeria listings; posting dates not consistently exposed."],
        ["Indeed Nigeria – Abuja DevOps roles", found, "DevOps, Senior DevOps, Linux Administrator", indeed_abuja, "Includes Abuja roles; verify location restrictions and active status."],
        ["Indeed Nigeria – remote DevOps roles surfaced for Lagos", found, "Remote DevOps Engineer", indeed_remote, "Remote results page; confirm Nigeria eligibility and direct application path."],
        ["Verification method", found, "Reviewed available job-board snippets and opened selected listings", "See links in Job Tracker", "Some job pages blocked direct retrieval; direct employer links were not consistently available."]
    ], columns=["Source", "Date checked", "Search focus", "URL", "Notes / limitations"]).to_excel(writer, sheet_name="Search Log", index=False)

wb = load_workbook("DevOps_Jobs_Nigeria_Lagos_12_Oct_2026.xlsx")
navy = "17365D"
blue = "1F4E78"
light_blue = "D9EAF7"
green = "E2F0D9"
yellow = "FFF2CC"
white = "FFFFFF"
thin_gray = Side(style="thin", color="D9E1F2")

for ws in wb.worksheets:
    ws.freeze_panes = "A2"
    ws.sheet_view.showGridLines = False
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 32
    for cell in ws[1]:
        cell.fill = PatternFill("solid", fgColor=navy)
        cell.font = Font(color=white, bold=True, size=10)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        cell.border = Border(bottom=Side(style="medium", color=blue))
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = Border(bottom=thin_gray)
            if cell.row % 2 == 0:
                cell.fill = PatternFill("solid", fgColor="F7FAFC")
    for col_idx in range(1, ws.max_column + 1):
        values = [str(ws.cell(row=r, column=col_idx).value or "") for r in range(1, ws.max_row + 1)]
        max_len = min(max(max((len(line) for line in v.splitlines()), default=0) for v in values) + 2, 48)
        ws.column_dimensions[get_column_letter(col_idx)].width = max(14, max_len)
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.page_margins.left = 0.25
    ws.page_margins.right = 0.25

ws = wb["Job Tracker"]
# Sensible widths for long text and link columns
widths = {
    "A": 12, "B": 34, "C": 28, "D": 30, "E": 20, "F": 28, "G": 25,
    "H": 55, "I": 42, "J": 42, "K": 42, "L": 38, "M": 34, "N": 52
}
for col, width in widths.items():
    ws.column_dimensions[col].width = width
for r in range(2, ws.max_row + 1):
    ws.row_dimensions[r].height = 82
    for c in (10, 11):
        cell = ws.cell(r, c)
        if cell.value and str(cell.value).startswith("http"):
            cell.hyperlink = cell.value
            cell.style = "Hyperlink"
    priority = ws.cell(r, 1)
    if priority.value == "High":
        priority.fill = PatternFill("solid", fgColor=green)
        priority.font = Font(bold=True, color="006100")
    elif priority.value == "Low":
        priority.fill = PatternFill("solid", fgColor=yellow)
    # Make links visibly clickable
    for c in (10, 11):
        cell = ws.cell(r, c)
        if cell.value and str(cell.value).startswith("http"):
            cell.hyperlink = str(cell.value)
            cell.font = Font(color="0563C1", underline="single")
            cell.alignment = Alignment(vertical="top", wrap_text=True)

# Add tracking status dropdown
from openpyxl.worksheet.datavalidation import DataValidation
status_col = ws.max_column + 1
ws.cell(1, status_col, "Application Status")
ws.cell(1, status_col).fill = PatternFill("solid", fgColor=navy)
ws.cell(1, status_col).font = Font(color=white, bold=True)
ws.cell(1, status_col).alignment = Alignment(vertical="center", wrap_text=True)
ws.column_dimensions[get_column_letter(status_col)].width = 20
for r in range(2, ws.max_row + 1):
    ws.cell(r, status_col, "To Apply")
    ws.cell(r, status_col).alignment = Alignment(vertical="top", wrap_text=True)
dv = DataValidation(type="list", formula1='"Not Applied,To Apply,Applied,Interview,Rejected,Closed"', allow_blank=False)
dv.error = "Choose a status from the dropdown."
dv.errorTitle = "Invalid status"
dv.prompt = "Track your application progress."
dv.promptTitle = "Application status"
ws.add_data_validation(dv)
dv.add(f"{get_column_letter(status_col)}2:{get_column_letter(status_col)}{ws.max_row}")

# Table on job tracker
ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"
tab = Table(displayName="DevOpsJobTracker", ref=ref)
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showFirstColumn=False, showLastColumn=False, showRowStripes=True, showColumnStripes=False)
ws.add_table(tab)

# Summary formatting
sumws = wb["Summary"]
sumws.column_dimensions["A"].width = 30
sumws.column_dimensions["B"].width = 115
for r in range(2, sumws.max_row + 1):
    sumws.row_dimensions[r].height = 55 if r < 9 else 240
    sumws.cell(r, 2).alignment = Alignment(vertical="top", wrap_text=True)
# Search log formatting
logws = wb["Search Log"]
for col, width in {"A": 42, "B": 18, "C": 45, "D": 70, "E": 80}.items():
    logws.column_dimensions[col].width = width
for r in range(2, logws.max_row + 1):
    logws.row_dimensions[r].height = 62
    cell = logws.cell(r, 4)
    if cell.value and str(cell.value).startswith("http"):
        cell.hyperlink = cell.value
        cell.font = Font(color="0563C1", underline="single")
        cell.alignment = Alignment(vertical="top", wrap_text=True)

wb.save("C:/Users/USER/Desktop/DevOps_Jobs_Nigeria_Lagos_12_Oct_2026.xlsx")
print("Workbook created with", len(jobs), "job leads and 3 sheets.")
