from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION

doc = Document()
sec = doc.sections[0]
sec.top_margin = Pt(30)
sec.bottom_margin = Pt(30)
sec.left_margin = Pt(40)
sec.right_margin = Pt(40)

styles = doc.styles
styles["Normal"].font.name = "Arial"
styles["Normal"].font.size = Pt(9.5)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("SHOBO TAIWO ADEFOWOPE")
r.bold = True
r.font.size = Pt(17)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p.add_run("DEVOPS ENGINEER | AWS | DOCKER | TERRAFORM | CI/CD")
r.bold = True
r.font.size = Pt(10)

p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("shoboadefowope2007@gmail.com  |  github.com/Taiwo17/")

def heading(text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(5)
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(text)
    r.bold = True
    r.font.size = Pt(11)
    return p

def bullet(text):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(1)
    p.add_run(text)

heading("PROFESSIONAL SUMMARY")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(2)
p.add_run(
    "DevOps Engineer with 3 years of hands-on experience across cloud infrastructure, "
    "containerization, CI/CD automation, Infrastructure as Code, Linux administration, "
    "and application deployment. Experienced in building automated delivery pipelines, "
    "containerizing Node.js applications, provisioning AWS infrastructure with Terraform, "
    "and deploying workloads to EC2 using GitHub Actions, GitHub OIDC, and AWS Systems Manager. "
    "Strong practical focus on automation, troubleshooting, secure deployments, and repeatable infrastructure."
)

heading("CORE TECHNICAL SKILLS")
skills = [
    ("Cloud & AWS", "AWS EC2, IAM, VPC, Security Groups, AWS Systems Manager (SSM), GitHub OIDC"),
    ("Infrastructure as Code", "Terraform, reusable infrastructure configuration, AWS provisioning"),
    ("Containers", "Docker, Docker Compose, Docker Hub, image tagging, health checks, container networking"),
    ("CI/CD", "GitHub Actions, automated testing, build and deployment workflows, release automation"),
    ("Linux & Scripting", "Linux administration, SSH, Bash, Python fundamentals, troubleshooting"),
    ("Version Control", "Git, GitHub, branching, pull requests, repository workflows"),
    ("Application & Databases", "Node.js, Express.js, REST APIs, MongoDB"),
    ("Deployment & Operations", "EC2 deployments, SSM-based deployments, environment variables, connectivity troubleshooting")
]
for k, v in skills:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{k}: ")
    r.bold = True
    p.add_run(v)

heading("PROFESSIONAL EXPERIENCE")
p = doc.add_paragraph()
p.paragraph_format.space_after = Pt(1)
r = p.add_run("DevOps Engineer | Project-Based / Contract Experience")
r.bold = True
p.add_run("  | 2023 – Present")
for t in [
    "Designed and maintained CI/CD workflows for application testing, container image publishing, and automated cloud deployment using GitHub Actions.",
    "Provisioned AWS compute, networking, IAM, security groups, and supporting infrastructure using Terraform and Infrastructure as Code practices.",
    "Containerized Node.js and Express applications with Docker and Docker Compose, including environment configuration, health checks, and service dependencies.",
    "Implemented GitHub OIDC for short-lived AWS authentication and reduced reliance on long-lived cloud credentials.",
    "Deployed applications to AWS EC2 using AWS Systems Manager and automated deployment workflows.",
    "Troubleshot Linux, SSH, Docker, cloud-init, deployment, port-conflict, and application health-check issues.",
    "Worked with Git/GitHub workflows and documented infrastructure and deployment processes for repeatable environments."
]:
    bullet(t)

heading("SELECTED PROJECTS")
p = doc.add_paragraph()
r = p.add_run("Blog API | AWS Deployment & CI/CD")
r.bold = True
p.add_run("  | github.com/Taiwo17/Blog-API-YT")
for t in [
    "Containerized a Node.js/Express API and MongoDB with Docker Compose, including service health checks and local/EC2 verification.",
    "Provisioned EC2, security groups, IAM roles, and related AWS infrastructure with Terraform.",
    "Built GitHub Actions automation to run tests, publish commit-tagged Docker images, and deploy to EC2 through AWS Systems Manager.",
    "Configured GitHub OIDC for temporary AWS credentials and restricted SSH access to an administrator IP.",
    "Diagnosed cloud-init, SSH connectivity, and deployment health-check issues; implemented bounded retries to improve startup reliability."
]:
    bullet(t)

p = doc.add_paragraph()
r = p.add_run("Trosc Backend | Containerization & CI/CD")
r.bold = True
p.add_run("  | github.com/Taiwo17/trosc-backend-devops")
for t in [
    "Forked an existing backend, created a containerization branch, and built a local Docker image.",
    "Ran the application with environment variables, resolved host port conflicts, and created a Node.js test workflow in GitHub Actions.",
    "Investigated Docker Compose image-pull and Docker Hub connectivity failures."
]:
    bullet(t)

heading("EDUCATION")
p = doc.add_paragraph()
r = p.add_run("MSc Economics, Part-time")
r.bold = True
p.add_run(" | University of Lagos, Nigeria | In progress, expected 2027")
p = doc.add_paragraph()
r = p.add_run("B.Sc. (Ed) Economics and Education")
r.bold = True
p.add_run(" | Olabisi Onabanjo University, Nigeria | 2019 | Second Class Upper (CGPA: 3.75)")

path = "C:/Users/USER/Desktop/Shobo_Taiwo_Adefowope_DevOps_Engineer_CV.docx"
doc.save(path)
print(path)
